"""Tests for Application Gate."""

import pytest
import tempfile
import os
from datetime import datetime
from app.storage.models import Job, JobStatus
from app.applications.gate import ApplicationGate, ApplicationGateResult, AmbiguityResolver


@pytest.fixture
def temp_data_dir_with_history():
    """Create a temporary data directory with application history."""
    with tempfile.TemporaryDirectory() as tmpdir:
        # Create a sample job_applications_master.xlsx
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws['A1'] = "Job ID"
            ws['B1'] = "Company"
            ws['C1'] = "Job Title"
            ws['D1'] = "Applied Date"

            # Add some applications
            ws['A2'] = "JOB_WIX_001"
            ws['B2'] = "Wix"
            ws['C2'] = "Project Manager"
            ws['D2'] = datetime(2024, 1, 15)

            ws['A3'] = "JOB_JFROG_001"
            ws['B3'] = "JFrog"
            ws['C3'] = "Technical Program Manager"
            ws['D3'] = datetime(2024, 1, 20)

            ws['A4'] = None  # No Job ID for this one
            ws['B4'] = "SolarWinds"
            ws['C4'] = "Release Manager"
            ws['D4'] = datetime(2024, 2, 10)

            wb.save(os.path.join(tmpdir, "job_applications_master.xlsx"))
        except ImportError:
            pass  # openpyxl not available, skip

        yield tmpdir


class TestApplicationGate:
    """Test Application Gate checking logic."""

    def test_new_job_passes_through(self, temp_data_dir_with_history):
        """A job not in history should pass through unchanged."""
        gate = ApplicationGate(temp_data_dir_with_history)

        job = Job(
            job_id="job_001",
            title="Unknown Role",
            company="Unknown Company",
            status=JobStatus.NEW
        )

        checked_job, was_applied = gate.check(job)

        assert was_applied is False
        assert checked_job.status == JobStatus.NEW  # Status unchanged

    def test_match_by_job_id(self, temp_data_dir_with_history):
        """Job ID match should mark as already applied."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            job = Job(
                job_id="job_001",
                title="Project Manager",
                company="Wix",
                job_req_id="JOB_WIX_001",
                status=JobStatus.NEW
            )

            checked_job, was_applied = gate.check(job)

            assert was_applied is True
            assert checked_job.status == JobStatus.ALREADY_APPLIED
            assert checked_job.application_note == "אין פעולה — לא להגיש שוב"
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_match_by_company_title(self, temp_data_dir_with_history):
        """Company + Title match should mark as already applied."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            job = Job(
                job_id="job_002",
                title="Technical Program Manager",
                company="JFrog",
                status=JobStatus.NEW
            )

            checked_job, was_applied = gate.check(job)

            assert was_applied is True
            assert checked_job.status == JobStatus.ALREADY_APPLIED
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_case_insensitive_company_match(self, temp_data_dir_with_history):
        """Company name matching should be case-insensitive."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            job = Job(
                job_id="job_003",
                title="TECHNICAL PROGRAM MANAGER",  # Different case
                company="JFROG",  # Different case
                status=JobStatus.NEW
            )

            checked_job, was_applied = gate.check(job)

            assert was_applied is True
            assert checked_job.status == JobStatus.ALREADY_APPLIED
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_no_match_different_company(self, temp_data_dir_with_history):
        """Same title at different company should NOT match."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            job = Job(
                job_id="job_004",
                title="Project Manager",
                company="Different Company",
                status=JobStatus.NEW
            )

            checked_job, was_applied = gate.check(job)

            assert was_applied is False
            assert checked_job.status == JobStatus.NEW
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_no_match_different_title(self, temp_data_dir_with_history):
        """Same company with different title should NOT match."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            job = Job(
                job_id="job_005",
                title="Different Title",
                company="Wix",
                status=JobStatus.NEW
            )

            checked_job, was_applied = gate.check(job)

            assert was_applied is False
            assert checked_job.status == JobStatus.NEW
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_job_id_preferred_over_title(self, temp_data_dir_with_history):
        """Job ID should be checked first (even if title doesn't match)."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            job = Job(
                job_id="job_006",
                title="Wrong Title",
                company="Wrong Company",
                job_req_id="JOB_WIX_001",  # This ID exists in history
                status=JobStatus.NEW
            )

            checked_job, was_applied = gate.check(job)

            assert was_applied is True
            assert checked_job.status == JobStatus.ALREADY_APPLIED
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_is_applied_convenience_method(self, temp_data_dir_with_history):
        """Test is_applied() convenience method."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            applied_job = Job(
                job_id="job_007",
                title="Project Manager",
                company="Wix"
            )

            new_job = Job(
                job_id="job_008",
                title="Unknown",
                company="Unknown"
            )

            assert gate.is_applied(applied_job) is True
            assert gate.is_applied(new_job) is False
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_empty_history_all_new(self):
        """With empty history, all jobs should be new."""
        with tempfile.TemporaryDirectory() as tmpdir:
            gate = ApplicationGate(tmpdir)

            job = Job(
                job_id="job_009",
                title="Any Title",
                company="Any Company"
            )

            checked_job, was_applied = gate.check(job)

            assert was_applied is False
            assert checked_job.status == JobStatus.NEW


class TestApplicationGateResult:
    """Test ApplicationGateResult helper class."""

    def test_result_new_job(self):
        """Result for new job."""
        job = Job(
            job_id="job_001",
            title="PM",
            company="Wix",
            status=JobStatus.NEW
        )

        result = ApplicationGateResult(job, was_matched=False)

        assert result.was_matched is False
        assert "New application" in result.summary()

    def test_result_already_applied(self):
        """Result for already applied job."""
        job = Job(
            job_id="job_001",
            title="PM",
            company="Wix",
            status=JobStatus.ALREADY_APPLIED
        )

        result = ApplicationGateResult(
            job,
            was_matched=True,
            match_type="company_title",
            confidence=0.95
        )

        assert result.was_matched is True
        assert result.match_type == "company_title"
        assert result.confidence == 0.95
        assert "Already applied" in result.summary()
        assert "company + title" in result.summary()


class TestAmbiguityResolver:
    """Test AmbiguityResolver stub."""

    def test_resolve_not_implemented(self):
        """Stub should return None (inconclusive)."""
        job = Job(
            job_id="job_001",
            title="PM",
            company="Wix"
        )

        result = AmbiguityResolver.resolve(job)

        assert result is None

    def test_is_implemented_returns_false(self):
        """Stub should report as not implemented."""
        assert AmbiguityResolver.is_implemented() is False


class TestApplicationGateStatuses:
    """Test that Application Gate properly sets status."""

    def test_status_set_to_already_applied(self, temp_data_dir_with_history):
        """When matched, status should be ALREADY_APPLIED."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            job = Job(
                job_id="job_010",
                title="Project Manager",
                company="Wix",
                status=JobStatus.NEW
            )

            checked_job, _ = gate.check(job)

            assert checked_job.status == JobStatus.ALREADY_APPLIED
            assert checked_job.status.value == "already_applied"
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_status_unchanged_if_no_match(self, temp_data_dir_with_history):
        """If no match, status should remain unchanged."""
        try:
            import openpyxl
            gate = ApplicationGate(temp_data_dir_with_history)

            job = Job(
                job_id="job_011",
                title="Unknown",
                company="Unknown",
                status=JobStatus.TO_APPLY
            )

            checked_job, _ = gate.check(job)

            assert checked_job.status == JobStatus.TO_APPLY  # Unchanged
        except ImportError:
            pytest.skip("openpyxl not installed")
