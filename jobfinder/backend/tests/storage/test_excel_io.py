"""Tests for Excel I/O module."""

import pytest
import tempfile
import os
from datetime import datetime
from app.storage.excel_io import ApplicationHistory, SkillsProfile, get_data_dir, get_application_history, get_skills_profile


@pytest.fixture
def temp_data_dir():
    """Create a temporary data directory with test files."""
    with tempfile.TemporaryDirectory() as tmpdir:
        # Create a sample job_applications_master.xlsx
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws['A1'] = "Job ID"
            ws['B1'] = "Company"
            ws['C1'] = "Job Title"
            ws['D1'] = "Applied Date"
            ws['E1'] = "Source"
            ws['F1'] = "Notes"

            ws['A2'] = "JOB_001"
            ws['B2'] = "Wix"
            ws['C2'] = "Project Manager"
            ws['D2'] = datetime(2024, 1, 15)
            ws['E2'] = "Greenhouse"
            ws['F2'] = "Applied via careers page"

            ws['A3'] = "JOB_002"
            ws['B3'] = "JFrog"
            ws['C3'] = "Technical Program Manager"
            ws['D3'] = datetime(2024, 1, 20)
            ws['E3'] = "LinkedIn"

            wb.save(os.path.join(tmpdir, "job_applications_master.xlsx"))
        except ImportError:
            pass  # openpyxl not available, skip

        # Create a sample skills profile
        profile_content = """# Anat Aminof - Skills Profile

## Technical Skills
- Project Management
- Program Management
- Agile/Scrum
- Python
- JavaScript
- SQL

## Certifications
- Scrum Master (CSM)
- ITIL Foundation

## Experience
- 8 years in technical project management
"""
        with open(os.path.join(tmpdir, "Anat_Aminof_skills_profile.md"), 'w') as f:
            f.write(profile_content)

        yield tmpdir


def test_skills_profile_load(temp_data_dir):
    """Test loading skills profile."""
    profile = SkillsProfile(os.path.join(temp_data_dir, "Anat_Aminof_skills_profile.md"))

    assert len(profile.get_skills()) > 0
    assert "Project Management" in profile.get_skills()


def test_skills_profile_has_skill(temp_data_dir):
    """Test checking if skill exists in profile."""
    profile = SkillsProfile(os.path.join(temp_data_dir, "Anat_Aminof_skills_profile.md"))

    assert profile.has_skill("Project Management") is True
    assert profile.has_skill("project management") is True  # case-insensitive
    assert profile.has_skill("Go programming") is False


def test_skills_profile_missing_file():
    """Test loading non-existent skills profile."""
    profile = SkillsProfile("/nonexistent/path/Anat_Aminof_skills_profile.md")
    assert profile.get_skills() == []
    assert profile.has_skill("anything") is False


def test_application_history_load(temp_data_dir):
    """Test loading application history from Excel."""
    try:
        import openpyxl
        history = ApplicationHistory(os.path.join(temp_data_dir, "job_applications_master.xlsx"))
        apps = history.get_applications()

        assert len(apps) == 2
        assert apps[0]["job_id"] == "JOB_001"
        assert apps[0]["company"] == "Wix"
        assert apps[0]["job_title"] == "Project Manager"
        assert apps[1]["company"] == "JFrog"
    except ImportError:
        pytest.skip("openpyxl not installed")


def test_application_history_has_application_by_job_id(temp_data_dir):
    """Test checking application by job ID."""
    try:
        import openpyxl
        history = ApplicationHistory(os.path.join(temp_data_dir, "job_applications_master.xlsx"))

        assert history.has_application(job_id="JOB_001") is True
        assert history.has_application(job_id="JOB_999") is False
    except ImportError:
        pytest.skip("openpyxl not installed")


def test_application_history_has_application_by_company_title(temp_data_dir):
    """Test checking application by company and title."""
    try:
        import openpyxl
        history = ApplicationHistory(os.path.join(temp_data_dir, "job_applications_master.xlsx"))

        assert history.has_application(company="Wix", job_title="Project Manager") is True
        assert history.has_application(company="Wix", job_title="TPM") is False
        assert history.has_application(company="Apple", job_title="PM") is False
    except ImportError:
        pytest.skip("openpyxl not installed")


def test_application_history_missing_file():
    """Test loading non-existent application history file."""
    history = ApplicationHistory("/nonexistent/path/job_applications_master.xlsx")
    assert history.get_applications() == []
    assert history.has_application(job_id="anything") is False


def test_get_data_dir():
    """Test getting data directory."""
    # Should return default "." if DATA_DIR not set
    import os as os_module
    old_data_dir = os_module.environ.get("DATA_DIR")

    try:
        # Clear DATA_DIR
        if "DATA_DIR" in os_module.environ:
            del os_module.environ["DATA_DIR"]

        data_dir = get_data_dir()
        assert data_dir == "."
    finally:
        # Restore
        if old_data_dir:
            os_module.environ["DATA_DIR"] = old_data_dir


def test_get_application_history_helper(temp_data_dir):
    """Test helper function for getting application history."""
    try:
        import openpyxl
        history = get_application_history(temp_data_dir)
        assert len(history.get_applications()) > 0
    except ImportError:
        pytest.skip("openpyxl not installed")


def test_get_skills_profile_helper(temp_data_dir):
    """Test helper function for getting skills profile."""
    profile = get_skills_profile(temp_data_dir)
    assert len(profile.get_skills()) > 0
